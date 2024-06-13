import os
import csv
import pandas as pd
import pyodbc  # For Access databases
from docx import Document  # For .docx files
from datetime import datetime
from openpyxl import load_workbook  # For .xlsx files
from bs4 import BeautifulSoup  # For .html files


class DataObject:
    def __init__(self, data):
        self.data = data

    def apply_gpt4(self, column_name):
        if column_name in self.data.columns:
            self.data[column_name] = self.data[column_name].apply(process_with_gpt4)


def process_with_gpt4(user_input):
    # Placeholder for GPT-4 processing
    return user_input


def read_data(file_path):
    extension = os.path.splitext(file_path)[1]
    if extension == '.csv':
        return DataObject(pd.read_csv(file_path))
    elif extension == '.txt':
        with open(file_path, 'r') as f:
            return DataObject(pd.DataFrame({'text': f.readlines()}))
    elif extension == '.docx':
        doc = Document(file_path)
        return DataObject(pd.DataFrame({'text': [p.text for p in doc.paragraphs]}))
    elif extension == '.xlsx':
        return DataObject(pd.read_excel(file_path))
    elif extension == '.html':
        with open(file_path, 'r') as f:
            soup = BeautifulSoup(f, 'html.parser')
            return DataObject(pd.read_html(str(soup.find('table')))[0])
    elif extension == '.accdb':
        conn_str = (
            r'DRIVER={Microsoft Access Driver (*.mdb, *.accdb)};'
            f'DBQ={file_path};'
        )
        conn = pyodbc.connect(conn_str)
        sql = 'SELECT * FROM YourTable'  # Replace YourTable with the actual table name
        return DataObject(pd.read_sql(sql, conn))
    else:
        return None

def write_processed_data(data_object, output_folder, filename):
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    output_filename = f"{filename.split('.')[0]}_processed_{datetime.now().strftime('%Y%m%d')}.csv"
    output_path = os.path.join(output_folder, output_filename)
    data_object.data.to_csv(output_path, index=False)



def write_processed_data(data_object, output_folder, filename):
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    output_filename = f"{filename.split('.')[0]}_processed_{datetime.now().strftime('%Y%m%d')}.csv"
    output_path = os.path.join(output_folder, output_filename)
    data_object.data.to_csv(output_path, index=False)

if __name__ == "__main__":
    folder_path = "YourFolder"  # Replace with the actual folder path
    output_folder = os.path.join(folder_path, "Processed_Data")

    for filename in os.listdir(folder_path):
        file_path = os.path.join(folder_path, filename)
        data_object = read_data(file_path)
        if data_object:
            data_object.apply_gpt4('YourColumnName')  # Replace YourColumnName with the actual column name
            write_processed_data(data_object, output_folder, filename)

