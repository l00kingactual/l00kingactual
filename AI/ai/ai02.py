import os
import pandas as pd
from datetime import datetime


class DataObject:
    def __init__(self, data):
        self.data = data

    def apply_gpt4(self, column_name):
        if column_name in self.data.columns:
            self.data[column_name] = self.data[column_name].apply(process_with_gpt4)


import os
import pandas as pd
import json
import pyodbc
from docx import Document
from pptx import Presentation
from PyPDF2 import PdfFileReader
from openpyxl import Workbook
from bs4 import BeautifulSoup

def read_data(file_path):
    extension = os.path.splitext(file_path)[1]
    data_object = None

    if extension == '.py':
        with open(file_path, 'r') as f:
            data_object = {'script': f.read()}
    elif extension == '.csv':
        data_object = pd.read_csv(file_path).to_dict()
    elif extension == '.txt':
        with open(file_path, 'r') as f:
            data_object = {'text': f.readlines()}
    elif extension == '.json':
        with open(file_path, 'r') as f:
            data_object = json.load(f)
    elif extension == '.xlsx':
        data_object = pd.read_excel(file_path).to_dict()
    elif extension == '.sql':
        with open(file_path, 'r') as f:
            data_object = {'sql': f.read()}
    elif extension == '.docx':
        doc = Document(file_path)
        data_object = {'text': [p.text for p in doc.paragraphs]}
    elif extension == '.pptx':
        prs = Presentation(file_path)
        data_object = {'slides': [slide.title for slide in prs.slides]}
    elif extension == '.pdf':
        pdf = PdfFileReader(file_path)
        data_object = {'text': [pdf.getPage(i).extractText() for i in range(pdf.getNumPages())]}
    elif extension == '.html':
        with open(file_path, 'r') as f:
            soup = BeautifulSoup(f, 'html.parser')
            data_object = {'html': str(soup)}

    return data_object

def write_to_excel(data_objects, output_folder):
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    wb = Workbook()
    for file_name, data_object in data_objects.items():
        ws = wb.create_sheet(title=file_name[:30])
        for key, value in data_object.items():
            ws.append([key, value])

    output_path = os.path.join(output_folder, 'combined_data.xlsx')
    wb.save(output_path)

if __name__ == "__main__":
    root_folder = os.getcwd()
    output_folder = os.path.join(root_folder, "\\playground\\data")

    data_objects = {}
    for subdir, _, files in os.walk(root_folder):
        for filename in files:
            file_path = os.path.join(subdir, filename)
            data_object = read_data(file_path)
            if data_object:
                data_objects[filename] = data_object

    write_to_excel(data_objects, output_folder)


from datetime import datetime
import pandas as pd
import os

def write_processed_data(data_object, output_folder, filename):
    if not os.path.exists(output_folder):
        os.makedirs(output_folder)

    # Convert the data object to a Pandas DataFrame
    if isinstance(data_object, dict):
        df = pd.DataFrame.from_dict(data_object, orient='index').reset_index()
        df.columns = ['Key', 'Value']
    else:
        df = pd.DataFrame(data_object)

    # Generate the output filename and path
    output_filename = f"{filename.split('.')[0]}_processed_{datetime.now().strftime('%Y%m%d')}.csv"
    output_path = os.path.join(output_folder, output_filename)

    # Write the DataFrame to a CSV file
    df.to_csv(output_path, index=False)




def read_local(root_folder, output_folder):
    for subdir, _, files in os.walk(root_folder):
        for filename in files:
            file_path = os.path.join(subdir, filename)
            data_object = read_data(file_path)
            if data_object:
                data_object.apply_gpt4('YourColumnName')  # Replace YourColumnName with the actual column name
                write_processed_data(data_object, output_folder, filename)


import os

def read_local(root_folder, output_folder):
    data_objects = {}
    print(f"Starting data processing from root folder: {root_folder}")

    for subdir, _, files in os.walk(root_folder):
        print(f"Reading files from directory: {subdir}")

        for filename in files:
            file_path = os.path.join(subdir, filename)
            print(f"Reading file: {file_path}")

            data_object = read_data(file_path)
            if data_object:
                print(f"Processing data from file: {filename}")
                data_objects[filename] = data_object

                print(f"Writing processed data for file: {filename}")
                write_processed_data(data_object, output_folder, filename)

    print("Writing all data objects to Excel workbook.")
    write_to_excel(data_objects, output_folder)
    print("Data processing complete.")

if __name__ == "__main__":
    root_folder = os.getcwd()  # Current root directory
    output_folder = os.path.join(root_folder, "data\\ai_output\\")

    if not os.path.exists(output_folder):
        os.makedirs(output_folder)
        print(f"Created output folder: {output_folder}")

    read_local(root_folder, output_folder)

